from pathlib import Path
import re, json
from xml.sax.saxutils import escape
from PIL import Image, ImageDraw, ImageFont
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, PageBreak, KeepTogether
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / '.uiux' / 'assets'
ASSETS.mkdir(parents=True, exist_ok=True)
OUT = ROOT / 'output' / 'pdf'
OUT.mkdir(parents=True, exist_ok=True)
FONT = Path('C:/Windows/Fonts/segoeui.ttf')
BOLD = Path('C:/Windows/Fonts/segoeuib.ttf')
pdfmetrics.registerFont(TTFont('UI', str(FONT)))
pdfmetrics.registerFont(TTFont('UIBold', str(BOLD)))
pdfmetrics.registerFontFamily('UI', normal='UI', bold='UIBold', italic='UI', boldItalic='UIBold')

INK='#152536'; MUTED='#516174'; PRIMARY='#0F766E'; BG='#F7F9FC'; BORDER='#CBD5E1'
def font(n=25, bold=False): return ImageFont.truetype(str(BOLD if bold else FONT), n)
def surface(w,h):
    im=Image.new('RGB',(w,h),'white'); return im,ImageDraw.Draw(im)
def rect(d,xy,fill='white',radius=12,line=BORDER): d.rounded_rectangle(xy,radius=radius,fill=fill,outline=line,width=2)
def label(d,xy,text,n=25,bold=False,fill=INK): d.text(xy,text,font=font(n,bold),fill=fill)
def button(d,xy,text,primary=True,n=23):
    rect(d,xy,PRIMARY if primary else 'white',8,PRIMARY if primary else BORDER)
    b=d.textbbox((0,0),text,font=font(n,True)); x1,y1,x2,y2=xy
    d.text(((x1+x2-(b[2]-b[0]))/2,(y1+y2-(b[3]-b[1]))/2-5),text,font=font(n,True),fill='white' if primary else INK)
def chip(d,x,y,text,w,selected=False):
    rect(d,(x,y,x+w,y+43),'#CCFBF1' if selected else 'white',20,PRIMARY if selected else BORDER)
    label(d,(x+13,y+6),text,21,selected,PRIMARY if selected else MUTED)
def save(im,name): im.save(ASSETS/name); return ASSETS/name

im,d=surface(1200,560)
rect(d,(0,0,1198,558),BG,12)
rect(d,(0,0,208,558),'white',12)
label(d,(22,22),'CampusNext',28,True)
for i,t in enumerate(['Today','Explore','Tasks','Teams','Profile']):
    if i==0: rect(d,(14,83+i*61,193,134+i*61),'#CCFBF1',9,PRIMARY)
    label(d,(30,95+i*61),t,24,i==0)
label(d,(239,28),'Today',33,True)
rect(d,(560,21,1158,70),'white',8); label(d,(580,31),'Search or ask CampusNext',24,fill=MUTED)
rect(d,(239,93,786,209)); label(d,(261,109),'Your next action',22,True)
label(d,(261,139),'Finish the hackathon abstract',27,True)
label(d,(261,175),'Due tomorrow  |  Opportunity requirement',19,fill=MUTED)
rect(d,(811,93,1158,193),'#FFF7ED'); label(d,(835,111),'Personal streak   7 days',25,True)
label(d,(835,154),'Today counted  |  Freeze: 1',21,fill='#92400E')
label(d,(239,218),'Needs your attention',25,True)
rect(d,(239,264,786,389)); label(d,(261,280),'Submission deadline changed',25,True)
label(d,(261,318),'Review the effect on your tasks',22,fill=MUTED)
button(d,(566,339,765,379),'Review change',False,21)
rect(d,(811,219,1158,389)); label(d,(834,239),'Buddy streak   4 days',24,True)
label(d,(834,280),'Your action counted',22); label(d,(834,313),'Buddy contribution pending',21,fill=MUTED)
label(d,(239,418),'For your interests',25,True)
rect(d,(239,463,1158,539)); label(d,(260,482),'Design Jam  |  Cultural Club',24,True)
label(d,(750,488),'Snacks included',22,fill=PRIMARY)
save(im,'today.png')

im,d=surface(1200,660)
for x in [8,612]: rect(d,(x,3,x+575,655),BG,22)
label(d,(32,21),'Explore',33,True); label(d,(637,21),'Hackathon details',33,True)
rect(d,(30,72,557,122)); label(d,(46,85),'Search opportunities',23,fill=MUTED)
chip(d,30,143,'Food',110,True); chip(d,154,143,'Goodies',136,True); chip(d,305,143,'Filters',115)
label(d,(30,208),'Must include both food and goodies',23,True)
label(d,(30,244),'2 matching opportunities',21,fill=MUTED)
rect(d,(30,285,557,507)); label(d,(50,302),'Campus Hackathon',28,True)
label(d,(50,347),'Tech Club  |  Registration in 4 days',23,fill=MUTED)
label(d,(50,387),'Lunch included',23,fill=PRIMARY)
label(d,(50,423),'Kit for first 100 registrants',23)
label(d,(50,459),'Stock unconfirmed  |  Verified source',21,fill=MUTED)
button(d,(351,523,555,570),'View event',True)
label(d,(641,80),'Tech Club  |  Verified source',23,fill=PRIMARY)
label(d,(641,122),'Registration deadline: in 4 days',23,True)
rect(d,(637,168,1162,298)); label(d,(659,184),'Readiness',25,True)
label(d,(659,225),'Team member: missing',23); label(d,(659,263),'Abstract: incomplete',23)
rect(d,(637,324,1162,516)); label(d,(659,340),'Food and goodies',25,True)
label(d,(659,380),'Lunch included',23)
label(d,(659,418),'Welcome kit for first 100 registrants',23)
label(d,(659,456),'Stock not confirmed',22,fill=MUTED)
button(d,(637,542,1162,596),'Save opportunity',True,24)
label(d,(664,616),'Today   Explore   Tasks   Teams   Profile',20,fill=MUTED)
save(im,'explore_detail.png')

im,d=surface(1200,420)
rect(d,(2,2,1198,418),BG)
label(d,(25,17),'Preparation plan',31,True)
label(d,(27,66),'Suggested tasks  |  Review before accepting',23,fill=MUTED)
for i,(t,o,s) in enumerate([('Choose project problem','You','Ready'),('Draft abstract','Unassigned','Needs owner'),('Submit registration','You','Blocked by abstract')]):
    y=112+i*76;rect(d,(24,y,1175,y+64));label(d,(44,y+17),t,25,True)
    label(d,(600,y+18),o,23,fill=MUTED); label(d,(849,y+18),s,22,fill='#92400E' if i else PRIMARY)
button(d,(931,355,1176,403),'Accept plan',True)
label(d,(28,365),'Find a teammate if the opportunity requires one',22,fill=MUTED)
save(im,'plan.png')

im,d=surface(1200,480)
rect(d,(2,2,1198,478),BG)
label(d,(26,19),'Hackathon team',31,True)
label(d,(27,65),'Campus Hackathon  |  Project preparation',22,fill=MUTED)
rect(d,(25,111,569,393)); label(d,(47,129),'Accepted members',25,True)
label(d,(47,179),'You  |  Product design',25,True)
label(d,(47,219),'Ananya  |  Frontend development',24)
label(d,(47,274),'Open role: backend development',23)
button(d,(48,323,307,374),'Find collaborator',True,22)
rect(d,(597,111,1174,393)); label(d,(621,129),'Pending invitation',25,True)
label(d,(621,179),'Arjun  |  Backend development',24,True)
label(d,(621,221),'Invited - awaiting acceptance',22,fill=MUTED)
label(d,(621,259),'Not yet a team member',22,fill='#92400E')
button(d,(916,323,1153,374),'Cancel invitation',False,21)
rect(d,(25,412,1174,460)); label(d,(43,422),'Draft abstract  |  Owner: You  |  Accepted task',23,True)
save(im,'teams.png')

im,d=surface(1200,410)
rect(d,(2,2,1198,408),BG)
label(d,(28,21),'Campus Streaks',31,True)
rect(d,(28,79,574,338),'#FFF7ED'); label(d,(50,95),'Personal   7 credited days',28,True)
label(d,(50,142),'Today counted: abstract milestone',23)
for i,t in enumerate(['M','T','W','T','F','S','S']):
    x=52+i*67; color='#CCFBF1' if i<6 else 'white'
    rect(d,(x,183,x+50,247),color,8);label(d,(x+14,187),t,24,True)
    if i<6: d.line([(x+17,230),(x+23,237),(x+35,224)],fill=PRIMARY,width=3)
    else: label(d,(x+20,224),'-',19,True,fill=MUTED)
label(d,(51,273),'1 freeze available  |  Refreshes Monday',22,fill=MUTED)
rect(d,(600,79,1174,338)); label(d,(625,96),'Buddy   4 credited days',28,True)
label(d,(625,147),'You: counted today',23,fill=PRIMARY)
label(d,(625,187),'Buddy: contribution pending',23)
label(d,(625,228),'A protected day pauses the count',22,fill=MUTED)
button(d,(913,270,1155,321),'View rules',False,22)
label(d,(29,367),'Campus timezone: Asia/Kolkata  |  Sharing: private',22,fill=MUTED)
save(im,'streaks.png')

im,d=surface(1200,555)
rect(d,(2,2,1198,553),BG)
label(d,(24,18),'Review announcement',31,True)
label(d,(24,65),'Review queue  >  Campus Hackathon',22,fill=MUTED)
rect(d,(25,111,504,482)); label(d,(47,127),'Original source',25,True)
rect(d,(47,181,480,391),'#F0F4F8'); label(d,(71,217),'CAMPUS HACKATHON',28,True)
label(d,(71,270),'Registration deadline and requirements',19)
label(d,(71,315),'Lunch included  |  Limited welcome kits',19)
label(d,(49,425),'Open source image',22,fill=PRIMARY)
rect(d,(529,111,1172,482));label(d,(553,129),'Extracted details',25,True)
for i,(t,v) in enumerate([('Title','Campus Hackathon'),('Deadline','Review extracted date'),('Eligibility','Not stated - needs review'),('Food','Provided - verify conditions'),('Goodies','First 100 registrants')]):
    y=177+i*55;label(d,(554,y),t,22,True);label(d,(735,y),v,21,fill='#92400E' if i in [1,2] else MUTED)
button(d,(771,496,956,540),'Save draft',False,21)
rect(d,(976,496,1172,540),'#E2E8F0',8);label(d,(1033,504),'Publish',21,True,fill=MUTED)
label(d,(28,506),'Verify the deadline before publishing',22,fill=MUTED)
save(im,'review.png')

SOURCE=(ROOT/'CampusNext_UIUX.md').read_text(encoding='utf-8')
sections=re.split(r'^## ',SOURCE,flags=re.M)
images={4:'today.png',5:'explore_detail.png',6:'plan.png',7:'teams.png',8:'streaks.png',9:'review.png'}
normal=ParagraphStyle('Body',fontName='UI',fontSize=10.8,leading=14.4,spaceAfter=9,textColor=colors.black)
heading=ParagraphStyle('Heading',fontName='UIBold',fontSize=18,leading=23,spaceAfter=13,textColor=colors.black)
title=ParagraphStyle('Title',fontName='UIBold',fontSize=26,leading=31,spaceAfter=13,textColor=colors.black)
cell=ParagraphStyle('Cell',parent=normal,fontSize=9.4,leading=12,spaceAfter=0)
cellh=ParagraphStyle('CellHeader',parent=cell,fontName='UIBold',textColor=colors.white)
caption=ParagraphStyle('Caption',parent=normal,fontSize=9,leading=12,textColor=colors.HexColor(MUTED),spaceAfter=12)
def rich(t):
    t=escape(t)
    return re.sub(r'\*\*(.*?)\*\*',r'<b>\1</b>',t)
def chunks(lines):
    i=0
    while i<len(lines):
        if not lines[i].strip():i+=1;continue
        if lines[i].startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                cols=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[-: ]+',x) for x in cols):rows.append(cols)
                i+=1
            yield 'table',rows
        elif lines[i].startswith('- '):
            yield 'bullet',lines[i][2:];i+=1
        else:
            p=[]
            while i<len(lines) and lines[i].strip() and not lines[i].startswith('|') and not lines[i].startswith('- '):p.append(lines[i]);i+=1
            yield 'p',' '.join(p)

story=[]
doc=Document(); sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
for name,size in [('Normal',11),('Title',26),('Heading 1',18),('Heading 2',14)]:
    st=doc.styles[name];st.font.name='Segoe UI';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_after=Pt(8);st.paragraph_format.line_spacing=1.15
doc.add_paragraph('CampusNext UI UX Design Specification','Title')
doc.add_paragraph('Version 1.0 | 2 October 2026')
story += [Paragraph('CampusNext UI UX Design Specification',title),Paragraph('Version 1.0 | 2 October 2026',caption)]

for index,section in enumerate(sections[1:],1):
    if index>1:story.append(PageBreak());doc.add_page_break()
    lines=section.splitlines(); h=lines[0].strip();story.append(Paragraph(h,heading));doc.add_paragraph(h,'Heading 1')
    if index in images:
        file=images[index]
        width=(6.5 if index==8 else 7.05)*72
        with Image.open(ASSETS/file) as source_image:
            height=width*source_image.height/source_image.width
        story.append(RLImage(str(ASSETS/file),width=width,height=height))
        story.append(Paragraph('Layout wireframe showing component hierarchy',caption))
        doc.add_picture(str(ASSETS/file),width=Inches(7.05));doc.add_paragraph('Layout wireframe showing component hierarchy')
    for kind,data in chunks(lines[1:]):
        if kind=='table':
            n=len(data[0]); widths=[110,200,197] if n==3 else [165,342]
            arr=[[Paragraph(rich(v),cellh if r==0 else cell) for v in row] for r,row in enumerate(data)]
            table=Table(arr,colWidths=widths,repeatRows=1,hAlign='LEFT')
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#16324F')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F3F6FA')]),('GRID',(0,0),(-1,-1),.6,colors.HexColor('#D9D9D9')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),9),('RIGHTPADDING',(0,0),(-1,-1),9),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
            story += [table,Spacer(1,12)]
            dt=doc.add_table(rows=1, cols=n);dt.autofit=False;dt.alignment=WD_TABLE_ALIGNMENT.LEFT
            for rowno,row in enumerate(data):
                cells=dt.rows[0].cells if rowno==0 else dt.add_row().cells
                for c,(value,width) in enumerate(zip(row,widths)):
                    cells[c].width=Inches(width/72);cells[c].text=value.replace('**','');cells[c].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    tcpr=cells[c]._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'16324F' if rowno==0 else 'F3F6FA' if rowno%2==0 else 'FFFFFF');tcpr.append(shd)
                    borders=OxmlElement('w:tcBorders')
                    for edge in ['top','left','bottom','right']:
                        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                    tcpr.append(borders)
                    for p in cells[c].paragraphs:
                        for run in p.runs:run.font.size=Pt(9.5);run.font.bold=rowno==0;run.font.color.rgb=RGBColor(255,255,255) if rowno==0 else RGBColor(0,0,0)
                if rowno==0:
                    prop=dt.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');prop.append(repeat)
            doc.add_paragraph()
        else:
            story.append(Paragraph(('&#8226; ' if kind=='bullet' else '')+rich(data),normal))
            p=doc.add_paragraph(style='List Bullet' if kind=='bullet' else 'Normal')
            for k,part in enumerate(re.split(r'(\*\*.*?\*\*)',data)):
                r=p.add_run(part[2:-2] if part.startswith('**') else part);r.bold=part.startswith('**')

def footer(c,d):
    c.setFont('UI',8);c.setFillColor(colors.HexColor(MUTED));c.drawString(50,26,'CampusNext UI UX specification');c.drawRightString(562,26,str(d.page))

pdf=OUT/'CampusNext_UIUX_Specification.pdf'
SimpleDocTemplate(str(pdf),pagesize=letter,rightMargin=50,leftMargin=50,topMargin=43,bottomMargin=43,title='CampusNext UI UX Design Specification',author='CampusNext project').build(story,onFirstPage=footer,onLaterPages=footer)
docx=ROOT/'.uiux'/'CampusNext_UIUX_Specification.docx';doc.save(docx)
print(json.dumps({'pdf':str(pdf),'word_source':str(docx),'wireframes':len(images),'sections':len(sections)-1}))
